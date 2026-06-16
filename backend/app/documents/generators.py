"""Document generators (Phase 3): Markdown / DOCX / PDF / XLSX.

Markdown always works (pure stdlib). DOCX/PDF/XLSX use optional libraries
(python-docx / reportlab / openpyxl), imported lazily, so the backend boots and
the core tests run without them. Each generator returns the written path.

A `Section`/`Document` model lets callers describe content once and render to
any format.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class Section:
    heading: str
    body: str = ""
    bullets: list[str] = field(default_factory=list)


@dataclass
class DocumentModel:
    title: str
    sections: list[Section] = field(default_factory=list)
    # Optional tabular data for spreadsheets: list of header + rows.
    table_headers: list[str] = field(default_factory=list)
    table_rows: list[list] = field(default_factory=list)


def render_markdown(doc: DocumentModel) -> str:
    lines = [f"# {doc.title}", ""]
    for s in doc.sections:
        lines.append(f"## {s.heading}")
        if s.body:
            lines.append(s.body)
        for b in s.bullets:
            lines.append(f"- {b}")
        lines.append("")
    if doc.table_headers:
        lines.append("| " + " | ".join(doc.table_headers) + " |")
        lines.append("| " + " | ".join("---" for _ in doc.table_headers) + " |")
        for row in doc.table_rows:
            lines.append("| " + " | ".join(str(c) for c in row) + " |")
        lines.append("")
    return "\n".join(lines)


def write_markdown(doc: DocumentModel, path: str) -> str:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(render_markdown(doc), encoding="utf-8")
    return path


def write_docx(doc: DocumentModel, path: str) -> str:
    try:
        from docx import Document as Docx
    except ImportError as exc:  # noqa: BLE001
        raise RuntimeError("`pip install python-docx` to generate DOCX.") from exc
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    d = Docx()
    d.add_heading(doc.title, level=0)
    for s in doc.sections:
        d.add_heading(s.heading, level=1)
        if s.body:
            d.add_paragraph(s.body)
        for b in s.bullets:
            d.add_paragraph(b, style="List Bullet")
    if doc.table_headers:
        table = d.add_table(rows=1, cols=len(doc.table_headers))
        for i, h in enumerate(doc.table_headers):
            table.rows[0].cells[i].text = str(h)
        for row in doc.table_rows:
            cells = table.add_row().cells
            for i, c in enumerate(row):
                cells[i].text = str(c)
    d.save(path)
    return path


def write_pdf(doc: DocumentModel, path: str) -> str:
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:  # noqa: BLE001
        raise RuntimeError("`pip install reportlab` to generate PDF.") from exc
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    styles = getSampleStyleSheet()
    story = [Paragraph(doc.title, styles["Title"]), Spacer(1, 12)]
    for s in doc.sections:
        story.append(Paragraph(s.heading, styles["Heading1"]))
        if s.body:
            story.append(Paragraph(s.body, styles["BodyText"]))
        if s.bullets:
            story.append(ListFlowable(
                [ListItem(Paragraph(b, styles["BodyText"])) for b in s.bullets], bulletType="bullet"
            ))
        story.append(Spacer(1, 8))
    SimpleDocTemplate(path, pagesize=letter).build(story)
    return path


def write_xlsx(doc: DocumentModel, path: str) -> str:
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # noqa: BLE001
        raise RuntimeError("`pip install openpyxl` to generate XLSX.") from exc
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = (doc.title or "Sheet")[:31]
    if doc.table_headers:
        ws.append(doc.table_headers)
        for row in doc.table_rows:
            ws.append(list(row))
    else:
        ws.append([doc.title])
        for s in doc.sections:
            ws.append([s.heading, s.body])
    wb.save(path)
    return path


# Format dispatch table.
WRITERS = {
    "md": write_markdown,
    "markdown": write_markdown,
    "docx": write_docx,
    "pdf": write_pdf,
    "xlsx": write_xlsx,
}


def generate(doc: DocumentModel, fmt: str, path: str) -> str:
    writer = WRITERS.get(fmt.lower())
    if writer is None:
        raise ValueError(f"Unsupported format: {fmt}")
    return writer(doc, path)
